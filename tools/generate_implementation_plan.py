from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"D:\Workshop\KhanShop\KhanShop Software Implementation Plan.docx")


PHASES = [
    (
        "Phase 0 — Requirements and Project Control",
        [
            "Convert the software blueprint into a prioritized implementation backlog.",
            "Separate MVP requirements from later enhancements.",
            "Define acceptance criteria for every module.",
            "Establish Git workflow, pull-request rules, coding conventions, and issue tracking.",
            "Prepare development, test, staging, and production environments.",
            "Record the non-negotiable business invariants listed below.",
        ],
        [
            "Stock cannot change without a corresponding stock transaction.",
            "Completed financial documents must not be physically deleted.",
            "Each sale must retain the cost price used at the time of sale.",
            "Returns and stock adjustments require the appropriate authorization.",
            "Customer and supplier balances must be derived from their ledgers.",
            "Online reservations and available stock must remain consistent.",
        ],
        "An approved backlog, acceptance criteria, and implementation roadmap are available.",
    ),
    (
        "Phase 1 — Solution Foundation",
        [
            "Create the .NET modular-monolith solution and its Api, Application, Domain, Infrastructure, Reporting, Shared, and Tests projects.",
            "Create the Angular application and feature-module structure.",
            "Configure PostgreSQL, EF Core, migrations, and environment-specific settings.",
            "Implement the base entity, audit fields, soft deletion, validation pipeline, global exception handling, and consistent API responses.",
            "Add pagination, filtering, UTC date handling, structured logging, database transactions, and a number-series generator.",
            "Create unit-test and integration-test projects.",
            "Configure continuous integration to build and test the backend and frontend.",
        ],
        [],
        "Both applications build successfully, migrations run, and health checks pass.",
    ),
    (
        "Phase 2 — Authentication and Authorization",
        [
            "Configure ASP.NET Core Identity.",
            "Implement JWT access tokens and refresh-token rotation.",
            "Add login, logout, refresh-token, forgot-password, reset-password, and current-user endpoints.",
            "Create Admin, Manager, and Salesperson roles.",
            "Implement permissions independently from roles and enforce permission-based policies.",
            "Seed the initial administrator, roles, and default permissions.",
            "Implement login history, activity logs, and the audit-log foundation.",
            "Create Angular login, session handling, route guards, and permission-aware UI directives.",
            "Add authentication and authorization tests.",
        ],
        [],
        "Every protected API and screen enforces the correct permissions.",
    ),
    (
        "Phase 3 — Shop Settings and Master Data",
        [
            "Implement shop, invoice, tax, payment-method, and system settings.",
            "Implement categories, subcategories, brands, product models, and units.",
            "Implement customer and supplier master data.",
            "Implement product variants, images, automatic product codes, and automatic barcodes.",
            "Support purchase price, sale price, minimum stock, warranty configuration, serial-number requirement, VAT applicability, and online availability.",
            "Add search, filters, paging, validation, soft deletion, and audit logging.",
            "Add barcode generation and printing.",
        ],
        [],
        "Master data and products can be managed without directly modifying stock.",
    ),
    (
        "Phase 4 — Inventory Foundation",
        [
            "Implement the stock transaction ledger and supported transaction types.",
            "Implement opening stock through stock transactions.",
            "Track available, reserved, damaged, warranty, and adjustment stock.",
            "Implement weighted-average costing.",
            "Implement stock-adjustment requests, approval, and rejection.",
            "Add current-stock, low-stock, damaged-stock, and stock-ledger screens.",
            "Add concurrency protection to prevent negative stock and overselling.",
            "Test every stock movement, reversal, approval, and balance calculation.",
        ],
        [],
        "Stock cannot be altered outside the inventory service, and every movement is traceable.",
    ),
    (
        "Phase 5 — Purchase and Supplier Accounting",
        [
            "Implement direct purchase invoices and purchase items.",
            "Support discounts, VAT, supplier invoice numbers, and changing purchase prices.",
            "Update weighted-average cost when a purchase is confirmed.",
            "Increase stock within the same database transaction as purchase confirmation.",
            "Implement full, partial, and mixed supplier payments.",
            "Create supplier payable and ledger entries.",
            "Implement purchase returns and corresponding stock reversals.",
            "Add purchase, payment, payable, return, and supplier-ledger screens.",
            "Test totals, stock effects, ledgers, concurrency, and transaction rollback.",
        ],
        [],
        "Purchases, supplier payments, balances, and inventory reconcile.",
    ),
    (
        "Phase 6 — POS Sales and Customer Accounting",
        [
            "Build a fast keyboard- and barcode-friendly POS screen.",
            "Add customer search, quick customer creation, product search, and barcode scanning.",
            "Implement item discounts, invoice discounts, optional VAT, and authorized price changes.",
            "Support cash, card, mobile banking, mixed payments, partial payment, and due sales.",
            "Save CostPriceAtSale for every sale line.",
            "Deduct stock atomically with invoice creation.",
            "Create customer-ledger and due entries.",
            "Generate A4 invoices and POS receipts.",
            "Add QR-based invoice and warranty verification.",
            "Test stock availability, totals, payments, printing, and concurrent sales.",
        ],
        [],
        "A sale updates inventory and customer accounting as one atomic operation.",
    ),
    (
        "Phase 7 — Quotations",
        [
            "Implement quotation creation, editing, cancellation, expiration, and printing.",
            "Support pricing, discounts, VAT, validity dates, and status history.",
            "Convert an approved quotation into a sale.",
            "Revalidate stock and current business rules during conversion.",
            "Ensure quotations never reserve or deduct stock.",
        ],
        [],
        "Quotation conversion produces a valid sale without changing stock prematurely.",
    ),
    (
        "Phase 8 — Customer Due Collection",
        [
            "Implement the customer ledger and balance calculation.",
            "Implement due collection using all supported payment methods.",
            "Define and implement the payment-allocation rule.",
            "Support return, refund, and authorized correction adjustments.",
            "Generate collection receipts.",
            "Add customer balance, statement, and payment-history reports.",
        ],
        [],
        "Each customer balance equals the sum of that customer's ledger entries.",
    ),
    (
        "Phase 9 — Sales Returns and Complaints",
        [
            "Require the original sales invoice for every return.",
            "Allow authorized users to submit return requests.",
            "Implement manager/admin approval and rejection.",
            "Record complaint reason and product condition.",
            "Support refund, replacement, customer-due adjustment, and rejection.",
            "Route products to available, damaged, warranty, or supplier-claim stock.",
            "Add approval history and audit logging.",
            "Prevent returns exceeding the quantity originally sold.",
            "Test partial returns and all financial and stock effects.",
        ],
        [],
        "Returned quantities, financial adjustments, and stock movements reconcile with the original sale.",
    ),
    (
        "Phase 10 — Serial Numbers and Warranty",
        [
            "Capture serial numbers during purchase or authorized stock entry.",
            "Require serial selection during the sale of serialized products.",
            "Calculate warranty start and expiry dates from the sale date.",
            "Implement warranty lookup by invoice or serial number.",
            "Implement claim creation, review, approval, rejection, supplier claim, repair, replacement, refund, and resolution.",
            "Maintain claim status history and notes.",
            "Prevent duplicate serial numbers and invalid warranty claims.",
            "Test warranty replacement and its stock movements.",
        ],
        [],
        "Every serialized unit has a traceable purchase, sale, and warranty history.",
    ),
    (
        "Phase 11 — Online Store and Orders",
        [
            "Build public product listing, filtering, search, and product-detail pages.",
            "Implement carts, checkout, delivery information, and delivery charges.",
            "Support website orders and manually entered Facebook, phone, and WhatsApp orders.",
            "Reserve stock when an order is confirmed.",
            "Deduct stock when an order is delivered or completed.",
            "Release reservations when an order is cancelled.",
            "Implement courier assignment, tracking numbers, delivery tracking, and status history.",
            "Create admin order-management screens.",
            "Reuse common payment, stock, ledger, and sales services where appropriate.",
            "Test simultaneous orders, cancellation, and reservation expiry.",
        ],
        [],
        "Online orders cannot oversell stock or leave incorrect reservations.",
    ),
    (
        "Phase 12 — Reporting and Dashboard",
        [
            "Implement reporting queries with Dapper.",
            "Add date ranges, paging, filters, and export support.",
            "Build sales, stock, valuation, purchase, customer due, supplier due, profit, return, warranty, product, and salesperson reports.",
            "Build dashboard cards, trends, top-selling products, and payment summaries.",
            "Restrict profit and sensitive reports with permissions.",
            "Reconcile report values against source transactions and ledgers.",
        ],
        [],
        "Stock, ledger, payment, and profit reports reconcile with transactional data.",
    ),
    (
        "Phase 13 — Excel Import and Export",
        [
            "Create downloadable templates for products, opening stock, customers, and suppliers.",
            "Implement upload validation, preview, and explicit confirmation.",
            "Track imports as controlled batches.",
            "Save valid rows and generate downloadable error files for invalid rows.",
            "Ensure opening-stock imports create stock transactions.",
            "Add export logging and permission checks.",
            "Test duplicate codes, duplicate barcodes, invalid references, and partial failures.",
        ],
        [],
        "Invalid rows cannot corrupt valid business data or bypass stock controls.",
    ),
    (
        "Phase 14 — Notifications and Background Jobs",
        [
            "Configure Hangfire and persistent job storage.",
            "Add low-stock alerts, daily sales summaries, due reminders, warranty reminders, order notifications, import cleanup, and large-report jobs.",
            "Initially generate copyable SMS and WhatsApp messages.",
            "Add external SMS or WhatsApp gateway integration as a later enhancement.",
            "Make jobs retry-safe and prevent duplicate notifications.",
            "Add job monitoring and failure alerts.",
        ],
        [],
        "Re-running a job does not duplicate financial, stock, or notification effects.",
    ),
    (
        "Phase 15 — Security, Reliability, and Performance",
        [
            "Review permission coverage for every API and screen.",
            "Add rate limiting to public and authentication endpoints.",
            "Validate uploads and restrict file types and sizes.",
            "Protect secrets and production configuration.",
            "Enforce HTTPS and appropriate security headers.",
            "Verify audit coverage for sensitive operations.",
            "Document database backup and restoration procedures.",
            "Run dependency, vulnerability, load, and concurrency tests.",
            "Optimize POS, checkout, stock, and reporting queries.",
        ],
        [],
        "Security review, recovery test, and performance targets pass.",
    ),
    (
        "Phase 16 — User Acceptance, Deployment, and Launch",
        [
            "Prepare representative test data and end-to-end business scenarios.",
            "Run user acceptance testing with Admin, Manager, and Salesperson roles.",
            "Reconcile stock, customer ledgers, supplier ledgers, and reports.",
            "Test A4 invoices, receipt printers, barcode scanners, and QR verification.",
            "Test desktop, tablet, and mobile layouts.",
            "Train users and prepare operational documentation.",
            "Deploy to staging, resolve acceptance findings, and obtain approval.",
            "Back up production before launch and prepare a rollback procedure.",
            "Deploy the production release.",
            "Monitor logs, jobs, performance, database growth, and business reconciliation after launch.",
        ],
        [],
        "The production system is accepted, backed up, monitored, and operational.",
    ),
]


MILESTONES = [
    ("1", "Foundation", "Phases 0–3", "Foundation, authentication, permissions, settings, and master data"),
    ("2", "Inventory & Purchasing", "Phases 4–5", "Controlled stock ledger, purchases, supplier payments, and payables"),
    ("3", "Retail Operations", "Phases 6–8", "POS, invoices, quotations, customer ledgers, and due collection"),
    ("4", "After-Sales Service", "Phases 9–10", "Returns, complaints, serial tracking, and warranties"),
    ("5", "Online Commerce", "Phase 11", "Public store, cart, orders, reservations, courier, and delivery"),
    ("6", "Management & Launch", "Phases 12–16", "Reports, imports, jobs, security, acceptance, and production launch"),
]


DEFINITION_OF_DONE = [
    "Approved acceptance criteria are satisfied.",
    "Backend API, Angular UI, and database changes are complete.",
    "A versioned database migration is included when required.",
    "Validation, permissions, audit logging, and soft-delete behavior are implemented.",
    "Unit, integration, and relevant end-to-end tests pass.",
    "Financial and stock side effects are covered by transaction and rollback tests.",
    "List screens include paging, search, and appropriate filters.",
    "Error handling and user feedback are clear.",
    "Documentation and changed-file summary are updated.",
    "The feature is demonstrated and accepted in the test environment.",
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def add_bullet(document, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = document.add_paragraph(text, style=style)
    paragraph.paragraph_format.space_after = Pt(2)
    return paragraph


def add_number(document, text):
    paragraph = document.add_paragraph(text, style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    return paragraph


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(28)
    styles["Title"].font.color.rgb = RGBColor(31, 78, 121)
    for name, size, color in [
        ("Heading 1", 18, RGBColor(31, 78, 121)),
        ("Heading 2", 14, RGBColor(46, 116, 181)),
        ("Heading 3", 11, RGBColor(68, 68, 68)),
    ]:
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = color

    header = section.header.paragraphs[0]
    header.text = "KhanShop  |  Software Implementation Plan"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(section.footer.paragraphs[0])

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("KhanShop Software\nImplementation Plan")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Retail Inventory, POS, Warranty & Online Order Management System"
    )
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(68, 68, 68)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Version 1.0  •  4 July 2026  •  Based on the Full Software Blueprint")
    meta.runs[0].font.size = Pt(9)
    meta.runs[0].font.color.rgb = RGBColor(110, 110, 110)

    document.add_paragraph()
    purpose = document.add_paragraph()
    purpose.alignment = WD_ALIGN_PARAGRAPH.CENTER
    purpose.add_run(
        "A practical, dependency-aware sequence for delivering the complete system "
        "as a tested modular monolith."
    ).italic = True

    document.add_page_break()

    document.add_heading("1. Purpose and Delivery Approach", level=1)
    document.add_paragraph(
        "This plan translates the Full Software Blueprint into an executable delivery "
        "sequence. The system will be built one vertical business module at a time, "
        "including its database model, backend API, Angular interface, permissions, "
        "audit trail, and automated tests."
    )
    document.add_paragraph(
        "The recommended architecture is a modular monolith using ASP.NET Core Web API "
        "on .NET 10, Angular, PostgreSQL, EF Core, Dapper for reporting, ASP.NET Core "
        "Identity with JWT, and Hangfire for background jobs."
    )

    document.add_heading("Guiding Principles", level=2)
    for item in [
        "Complete and verify each dependency before building the modules that rely on it.",
        "Deliver thin end-to-end slices instead of building the entire backend before the UI.",
        "Treat inventory and financial ledgers as controlled accounting records.",
        "Use database transactions for purchases, sales, returns, payments, and stock operations.",
        "Keep online orders distinct from POS sales while reusing shared business services.",
        "Use production-like test data and reconcile balances at every milestone.",
    ]:
        add_bullet(document, item)

    document.add_heading("2. Delivery Milestones", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Milestone", "Name", "Scope", "Business Outcome"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, "1F4E79")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for milestone, name, scope, outcome in MILESTONES:
        cells = table.add_row().cells
        for index, text in enumerate([milestone, name, scope, outcome]):
            cells[index].text = text

    document.add_page_break()
    document.add_heading("3. Detailed Implementation Plan", level=1)

    for phase_name, tasks, invariants, exit_condition in PHASES:
        document.add_heading(phase_name, level=2)
        document.add_heading("Implementation steps", level=3)
        for task in tasks:
            add_number(document, task)
        if invariants:
            document.add_heading("Business invariants", level=3)
            for invariant in invariants:
                add_bullet(document, invariant)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(5)
        run = paragraph.add_run("Exit criterion: ")
        run.bold = True
        paragraph.add_run(exit_condition)

    document.add_page_break()
    document.add_heading("4. Definition of Done", level=1)
    document.add_paragraph(
        "A phase or feature is complete only when all applicable conditions below are met."
    )
    for item in DEFINITION_OF_DONE:
        add_bullet(document, item)

    document.add_heading("5. Recommended Execution Cycle", level=1)
    cycle = [
        "Review the relevant blueprint section and confirm acceptance criteria.",
        "Design the domain rules, database changes, API contract, permissions, and UI behavior.",
        "Implement the smallest usable vertical slice.",
        "Add migration, validation, audit logging, and automated tests.",
        "Run code review, integration tests, and business reconciliation.",
        "Demonstrate the feature in the test environment and collect feedback.",
        "Resolve findings, update documentation, and close the backlog item.",
    ]
    for item in cycle:
        add_number(document, item)

    document.add_heading("6. Suggested Release Gates", level=1)
    gates = [
        (
            "Architecture gate",
            "Foundation builds cleanly; dependency boundaries and conventions are documented.",
        ),
        (
            "Security gate",
            "Authentication, permission policies, audit logging, and secret handling are verified.",
        ),
        (
            "Inventory gate",
            "Every stock-changing scenario creates balanced and traceable stock transactions.",
        ),
        (
            "Financial gate",
            "Sales, purchases, payments, returns, and ledger balances reconcile.",
        ),
        (
            "Online-order gate",
            "Reservations, delivery deductions, and cancellations pass concurrency tests.",
        ),
        (
            "Production gate",
            "Acceptance, backup restoration, security review, printer testing, and rollback rehearsal pass.",
        ),
    ]
    gate_table = document.add_table(rows=1, cols=2)
    gate_table.style = "Table Grid"
    for index, text in enumerate(["Gate", "Required Evidence"]):
        cell = gate_table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, "1F4E79")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for gate, evidence in gates:
        cells = gate_table.add_row().cells
        cells[0].text = gate
        cells[1].text = evidence

    document.add_heading("7. Immediate Next Action", level=1)
    document.add_paragraph(
        "Begin with Phase 0 by converting the blueprint into a prioritized backlog. "
        "Then execute Phase 1 as the first development increment: create the solution "
        "foundation, Angular application, PostgreSQL configuration, shared infrastructure, "
        "test projects, and continuous-integration build."
    )

    document.core_properties.title = "KhanShop Software Implementation Plan"
    document.core_properties.subject = (
        "Step-by-step implementation plan for the KhanShop retail management system"
    )
    document.core_properties.author = "KhanShop Project Team"
    document.core_properties.keywords = (
        "KhanShop, implementation plan, retail, inventory, POS, warranty, online orders"
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
